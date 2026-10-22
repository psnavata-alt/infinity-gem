"""CGE blind test ("Test it yourself", test 2; 8 Oct 2026): YOUR shock, checked by Excel's own formulas.

The tester writes any shock into YOUR_SHOCK.xlsx (dropdowns; any sizes). This script:
  1. reads the shock, and solves Infinity GEM on the same 9-sector Philippine SAM used for the IFPRI comparison (long-run
     closure, as in the verification workbook);
  2. builds a copy of the Excel verification workbook containing the base run, IFPRI's four test shocks AND the
     tester's shock ("YOURS", selected on opening). In it, the published equations of the IFPRI standard model
     (Lofgren, Harris and Robinson 2002) are ordinary Excel formulas evaluated at Infinity GEM's solution.
If the solution satisfies every equation for a shock that nobody chose in advance, it cannot have been "rigged" to
reproduce known numbers: it is a solution of the published model.

Run:  py tools/blind_cge_test.py --make-form                    (writes the template form; done once)
      py tools/blind_cge_test.py --form <your YOUR_SHOCK.xlsx> [--out <folder>] [--no-open]
      (reads the form, builds My_shock_check_<time>.xlsx in --out, default: a "results" folder next to the form, and opens it)
Written 8 Oct 2026 with an AI coding assistant.
"""
from __future__ import annotations

import datetime as dt
import importlib.util
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "docs" / "verification" / "test_it_yourself"
FORM = KIT / "templates" / "YOUR_SHOCK.xlsx"   # the shared template; testers work on their own copy (--form)
SECTORS = {"Agriculture, forestry, fishing": "agri", "Mining": "mine", "Food, beverages": "food",
           "Chemicals and petroleum": "chem", "Other manufacturing": "oman", "Electricity, water": "util",
           "Construction": "cons", "Trade and transport": "trad", "Other services": "serv"}
LEVERS = {  # label: (shock key in Infinity GEM, prefix of the item code or None for economy-wide levers)
    "World import price of a good": ("pwm_scale", "c"),
    "World export price of a good": ("pwe_scale", "c"),
    "Import tariff rate on a good": ("tm_scale", "c"),
    "Real investment demand for a good": ("qinv_scale", "c"),
    "Productivity of an industry": ("tfp_scale", "a"),
    "Real government consumption (all goods)": ("qg_scale", None),
    "Remittances to all households": ("remit_scale", None),
    "Capital stock (economy-wide)": ("qfs_scale", "fcap"),
}


def make_form():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Your shock"
    ws["A1"] = "Test it yourself (2): write ANY shock you like. Save, close, then double-click 'RUN MY SHOCK.bat'."
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("One row per change. Pick the lever and the sector from the lists; type the change in percent "
                "(e.g. 37 for +37%, -13 for -13%). Economy-wide levers ignore the sector. Up to 10 rows.")
    ws["A2"].font = Font(italic=True, color="555555")
    ws.append([])
    ws.append(["Lever", "Sector (if the lever needs one)", "Change (%)"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
    ws.append(["World import price of a good", "Agriculture, forestry, fishing", 37])   # an example: replace it
    ws.append(["Remittances to all households", "", -13])
    dv1 = DataValidation(type="list", formula1='"' + ",".join(LEVERS) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(SECTORS) + '"', allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    dv1.add("A5:A14"); dv2.add("B5:B14")
    for r in range(5, 15):
        for c in "ABC":
            ws[f"{c}{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A16"] = "The two rows above are only an example (rice-type farm imports +37%, remittances -13%): replace them with your own."
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 12
    FORM.parent.mkdir(parents=True, exist_ok=True)
    wb.save(FORM)
    print("wrote", FORM)


def read_form(form):
    ws = openpyxl.load_workbook(form, data_only=True)["Your shock"]
    spec, text = {}, []
    for lever, sector, pct in ws.iter_rows(min_row=5, max_row=14, max_col=3, values_only=True):
        if not lever or pct in (None, ""):
            continue
        key, pre = LEVERS[str(lever).strip()]
        f = 1 + float(pct) / 100
        if f <= 0:
            raise ValueError(f"A change of {pct}% would make '{lever}' zero or negative.")
        if pre in ("c", "a"):
            if sector not in SECTORS:
                raise ValueError(f"Row '{lever}' needs a sector from the list.")
            item = pre + SECTORS[sector]
            spec.setdefault(key, {})[item] = spec.get(key, {}).get(item, 1.0) * f
            text.append(f"{lever}: {sector} {float(pct):+g}%")
        elif pre == "fcap":
            spec.setdefault(key, {})["fcap"] = f
            text.append(f"{lever} {float(pct):+g}%")
        else:
            spec[key] = spec.get(key, 1.0) * f
            text.append(f"{lever} {float(pct):+g}%")
    if not spec:
        raise ValueError("The form has no shock: fill in at least one row.")
    return spec, "; ".join(text)


def main():
    if "--make-form" in sys.argv:
        return make_form()
    arg = lambda k: sys.argv[sys.argv.index(k) + 1] if k in sys.argv else None
    form = Path(arg("--form") or FORM)
    if not form.exists():
        raise FileNotFoundError(f"Shock form not found: {form}")
    no_open = "--no-open" in sys.argv
    out = Path(arg("--out") or form.parent / "results")       # read before sys.argv is cleared below
    spec, text = read_form(form)
    spec_ = importlib.util.spec_from_file_location("bw", ROOT / "tools" / "build_verification_workbook.py")
    bw = importlib.util.module_from_spec(spec_)
    sys.argv = sys.argv[:1]
    spec_.loader.exec_module(bw)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    bw.OUT = out
    bw.SCEN = list(bw.SCEN) + [("YOURS", spec)]
    bw.SCEN_TEXT = dict(bw.SCEN_TEXT, YOURS="YOUR shock: " + text.replace('"', "'"))
    bw.DEFAULT_SCEN = "YOURS"
    bw.WB_NAME = f"My_shock_check_{stamp}.xlsx"
    print("Your shock:", text)
    print("Infinity GEM spec:", spec)
    bw.build()                                   # stops with an error if Infinity GEM cannot solve the shock
    path = out / bw.WB_NAME
    print("\nOpen", path, "in Excel: sheet 'Checks' shows your shock (cell C3 = YOURS). Every equation group must say PASS.")
    if os.name == "nt" and not no_open:
        os.startfile(path)


if __name__ == "__main__":
    main()
