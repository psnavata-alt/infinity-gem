"""Electricity's fuel mix in PSA's 2018 tables against DOE's generation mix (23 Oct 2026).

RESEARCH, NOT FOR CITATION. Reads only files already on disk (no downloads, nothing outward):
  PSA 2018 Supply and Use Tables and the 240-industry I-O transaction table, Table C.1 (data/raw/io/);
  DOE 2025 Power Statistics, summary table (data/raw/energy/DOE_power_statistics_2025_edition_downloaded_2026-10-18/):
     gross generation by plant type, 2003-2025 (2003-2020 include off-grid generation; from 2021 grid only, DOE notes 1, 3);
  GTAP 10A Data Base, reference year 2014 (the GTAP 10A archive (licensed; not distributed) Aguiar, Chepeliev, Corong,
     McDougall and van der Mensbrugghe 2019, JGEA 4(1)): the Philippine electricity sector's energy use in Mtoe (gsdvole.har,
     EDF + EIF, which GTAP builds from IEA energy balances) and its purchases in US$ million at agents' prices (basedata.har,
     VDFA + VIFA; factor payments EVFA). Extracted to a temporary folder outside this repository and deleted after reading;
  World Bank Pink Sheet Dubai crude and BIS pesos per US dollar (data/raw/macro/), annual averages of the monthly files;
  PSA GVA in manufacturing at constant 2018 prices (OpenSTAT 0462B5BMFA1, data/raw/macro/);
  the OECD ICIO electricity cost shares already derived by research/icio/oil_cost_push_2022.py
     (outputs/icio_oil_cost_push_20261023_1648_electricity_fuel_inputs.csv; the raw ICIO stays outside this repository).

What it shows (printed, and saved as CSV tables in outputs/):
  1. PSA's 240-industry I-O table is the Use table under the fixed product-sales structure: T = D U, with D[r, p] = industry
     r's share of the domestic production of product p (Supply table), applied to total flows, imports included (and the
     imports column = D x imports by product); verified to about 1e-9 (₱ million). So each I-O row is an INDUSTRY's row: the
     'refined petroleum' row carries the products PSA's Supply table credits to the refining industry, about 48% of all
     crude-oil-and-natural-gas flows and 27% of coal flows among them.
  2. The electricity industry's fuel products (Use table) and I-O rows; the coal product's users; the secondary products
     of the refining industry; the electricity produced outside the electricity industry.
  3. A DOE- and IEA-consistent fuel bill of the electricity industry in 2018 (and the oil part in 2023), in pesos:
       oil   = DOE oil-based generation x GTAP 2014's oil products per MWh of oil-based generation x GTAP 2014's unit value,
               moved by the Dubai crude price and the peso-dollar rate;
       gas   = PSA's own Use-table value (natural gas bought by the electricity industry, product 040), kept;
       coal  = PSA's fuel total less oil and gas (the residual), cross-checked against PSA's coal supply x the power sector's
               share of coal use (GTAP 2014's IEA-based shares moved to 2018 by DOE coal generation and PSA real GVA).
Outputs: outputs/electricity_fuel_mix_evidence_<stamp>_<table>.csv; data/processed/electricity_fuel_mix_inputs_2014-2023.csv
(the inputs tools/build_own_sam_2023.py reads for version '0.6e'; one row per number, with its source).
Run (Git Bash): PYTHONIOENCODING=utf-8 py research/electricity_fuel_mix/fuel_mix_evidence.py      (about 30 seconds)
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from infinity_gem import paths                                          # noqa: E402

IO_DIR = paths.DATA_RAW / "io"
SUT = IO_DIR / "2018 Supply and Use Tables.xlsx"
TT240 = IO_DIR / "240-Industry20Table202018.xlsx"
DOE_PDF = (paths.DATA_RAW / "energy" / "DOE_power_statistics_2025_edition_downloaded_2026-10-18" /
           "DOE_1_Summary_electricity_consumption_demand_generation_capacity_2003-2025.pdf")
GTAP_PKG = ROOT / "docs" / "references" / "cge_models" / "gtap" / "GTAP10A_GTAP_AY.pkg"
ICIO_CSV = paths.output_dir() / "icio_oil_cost_push_20261023_1648_electricity_fuel_inputs.csv"
INPUTS_CSV = paths.DATA_PROCESSED / "electricity_fuel_mix_inputs_2014-2023.csv"
ELEC, STEAM = 172, 173                                    # PSA 240-industry codes
COAL, CRUDE_GAS, REFINED, OTHER_PETR = 39, 40, 107, 108
BASIC_IRON_STEEL, CEMENT = 123, 121
MJ_PER_TOE = 41868.0
GJ_PER_TONNE_5000KCAL = 5000 * 4.1868 / 1000              # 20.93 GJ per tonne of 5,000 kcal/kg coal (plausibility only)
TBTU_PER_PJ = 0.947817


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


# ----------------------------------------------------------------------------------------------------------- PSA tables
def psa_sut() -> dict:
    """PSA 2018 Supply and Use Tables (240 products x 240 industries, ₱ million, purchasers' prices)."""
    wb = openpyxl.load_workbook(SUT, read_only=True, data_only=True)
    U = list(wb["240_Use"].iter_rows(values_only=True))
    S = list(wb["240_Supply"].iter_rows(values_only=True))
    hu = [str(h).strip().replace("\n", " ") if h is not None else "" for h in U[9]]
    hs = [str(h).strip().replace("\n", " ") if h is not None else "" for h in S[9]]
    assert str(U[9 + REFINED][4]).strip() == "Refined petroleum products", U[9 + REFINED][:5]
    assert hu[4 + ELEC].startswith("Electric power generation"), hu[4 + ELEC]
    use = np.array([[_num(U[9 + i][4 + j]) for j in range(1, 241)] for i in range(1, 241)])
    sup = np.array([[_num(S[9 + i][4 + j]) for j in range(1, 241)] for i in range(1, 241)])
    lab = {str(r[4]).strip(): r for r in U[250:256] if r[4] is not None}
    go = np.array([_num(lab["Gross Output"][4 + j]) for j in range(1, 241)])
    jm = hs.index("Total Imports")
    imp = np.array([_num(S[9 + i][jm]) for i in range(1, 241)])
    fd0 = hu.index("Total Intermediate Demand")
    fd = {h: np.array([_num(U[9 + i][k]) for i in range(1, 241)]) for k, h in enumerate(hu) if k > fd0 and h}
    return {"use": use, "supply": sup, "gross_output": go, "imports": imp, "fd": fd,
            "products": [str(U[9 + k][4]).strip() for k in range(1, 241)], "industries": [hu[4 + j] for j in range(1, 241)]}


def psa_tt() -> dict:
    """PSA 2018 240-industry I-O transaction table, Table C.1 (₱ million): flows and the imports column."""
    ws = openpyxl.load_workbook(TT240, read_only=True, data_only=True)["240 TT"]
    R = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in R[9]]
    T = np.array([[_num(R[9 + i][1 + j]) for j in range(1, 241)] for i in range(1, 241)])
    jt = hdr.index("Imports of Goods and Services")
    lab = {str(r[0]).strip(): r for r in R if r and r[0] is not None}
    ti = np.array([_num(lab["TI"][1 + j]) for j in range(1, 241)])
    return {"T": T, "imports": np.abs(np.array([_num(R[9 + i][jt]) for i in range(1, 241)])), "total_inputs": ti}


def market_shares(sup: np.ndarray) -> np.ndarray:
    """D[r, p]: industry r's share of the domestic production of product p (Supply table)."""
    q = sup.sum(axis=1)
    return (sup / np.where(q == 0, 1.0, q)[:, None]).T


# ------------------------------------------------------------------------------------------------------- DOE, GTAP, prices
def doe_generation() -> pd.DataFrame:
    """DOE 2025 Power Statistics: gross power generation by plant type, GWh, 2003-2025 (rows: plant types; columns: years)."""
    import pypdf
    r = pypdf.PdfReader(str(DOE_PDF))
    txt = "\n".join(p.extract_text(extraction_mode="layout") for p in r.pages)
    blk = txt[txt.index("Gross Power Generation by Plant Type"):txt.index("Installed Capacity by Plant Type")]
    rows = {}
    for line in blk.splitlines()[1:]:
        nums = re.findall(r"\d{1,3}(?:,\d{3})*", line)
        lab = re.sub(r"[\d,\s]+$", "", line).strip()
        if lab and len(nums) == 23:
            rows[lab] = [float(n.replace(",", "")) for n in nums]
    need = ("Coal", "Oil-Based", "Natural Gas", "Renewable Energy (RE)", "Total Gross Generation")
    miss = [k for k in need if k not in rows]
    if miss:
        raise ValueError(f"DOE rows not parsed: {miss}")
    d = pd.DataFrame(rows, index=range(2003, 2026)).T
    gap = (d.loc[["Coal", "Oil-Based", "Natural Gas", "Renewable Energy (RE)"]].sum() - d.loc["Total Gross Generation"]).abs().max()
    if gap > 3:                                                    # DOE's own rounding
        raise ValueError(f"DOE plant types do not add up to the total (gap {gap} GWh)")
    return d


def gtap_2014_power() -> dict:
    """GTAP 10A (2014), Philippines: the electricity sector's energy use (Mtoe) and purchases (US$ million, agents' prices),
    and coal use by sector (Mtoe)."""
    import harpy
    from harpy.har_file_io import HarFileIO
    tmp = Path(tempfile.mkdtemp(prefix="gtap10a_2014_"))
    try:
        z = zipfile.ZipFile(GTAP_PKG)
        for n in ("basedata.har", "gsdvole.har"):
            (tmp / n).write_bytes(z.read(f"GTAP10A/GTAP/2014/{n}"))
        hfi = HarFileIO.readHarFileInfo(str(tmp / "gsdvole.har"))
        EDF, EIF = (HarFileIO.readHeader(hfi=hfi, header_name=h) for h in ("EDF", "EIF"))
        E, P, R = ([str(x).strip() for x in EDF["sets"][k]["dim_desc"]] for k in range(3))
        r = R.index("phl")
        vol = np.asarray(EDF["array"])[:, :, r] + np.asarray(EIF["array"])[:, :, r]          # Mtoe, energy x sector
        bd = harpy.HarFileObj.loadFromDisk(str(tmp / "basedata.har"))
        V = bd.getHeaderArrayObj("VDFA")
        C, PP, RR = ([str(x).strip() for x in V["sets"][k]["dim_desc"]] for k in range(3))
        r2 = RR.index("phl")
        j = PP.index("ely")
        val = np.asarray(V["array"])[:, j, r2] + np.asarray(bd.getHeaderArrayObj("VIFA")["array"])[:, j, r2]
        evfa = np.asarray(bd.getHeaderArrayObj("EVFA")["array"])[:, j, r2]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    je = P.index("ely")
    v = lambda e: float(vol[E.index(e), je])
    u = lambda c: float(val[C.index(c)])
    coal = vol[E.index("coa")]
    out = {"ely coal Mtoe": v("coa"), "ely oil products Mtoe": v("oil") + v("p_c"), "ely gas Mtoe": v("gas") + v("gdt"),
           "ely coal US$m": u("coa"), "ely oil products US$m": u("oil") + u("p_c"), "ely gas US$m": u("gas") + u("gdt"),
           "ely total cost US$m": float(val.sum() + evfa.sum()),
           "coal use, all sectors Mtoe": float(coal.sum()), "coal use, nmm (cement etc.) Mtoe": float(coal[P.index("nmm")]),
           "coal use, i_s (iron and steel) Mtoe": float(coal[P.index("i_s")])}
    out["coal use, other sectors Mtoe"] = out["coal use, all sectors Mtoe"] - out["ely coal Mtoe"] - \
        out["coal use, nmm (cement etc.) Mtoe"] - out["coal use, i_s (iron and steel) Mtoe"]
    return out


def annual_mean(file: str, col: str, year: int) -> float:
    d = pd.read_csv(paths.DATA_RAW / "macro" / file, comment="#")
    return float(d.loc[d["month"].astype(str).str.startswith(str(year)), col].mean())


def dubai(year: int) -> float:
    return annual_mean("WorldBank_PinkSheet_crude_oil_monthly_1960-2026.csv", "crude_dubai_usd_bbl", year)


def php_per_usd(year: int) -> float:
    return annual_mean("BIS_WS_XRU_PHP_per_USD_monthly_average.csv", "value", year)


def real_gva_growth(label: str, y0: int, y1: int) -> float:
    """PSA GVA in manufacturing at constant 2018 prices (OpenSTAT 0462B5BMFA1), y1 / y0, for the line starting with label."""
    d = pd.read_csv(paths.DATA_RAW / "macro" / "PSA_OpenSTAT_0462B5BMFA1_GVA_manufacturing_by_industry_annual.csv", comment="#")
    name = d.iloc[:, 0].astype(str).str.lstrip(".").str.strip()
    r = d[name.str.startswith(label)].iloc[0]
    return float(r[f"At Constant 2018 Prices {y1}"]) / float(r[f"At Constant 2018 Prices {y0}"])


# ------------------------------------------------------------------------------------------------------------ estimates
def oil_products_bill(gen_oil_gwh: float, year: int, g: dict, gen_oil_2014: float) -> dict:
    """Oil products bought by power plants, ₱ bn: generation x GTAP 2014 oil per MWh x GTAP 2014 unit value moved by Dubai
    crude and the peso-dollar rate."""
    toe_mwh = g["ely oil products Mtoe"] * 1e6 / (gen_oil_2014 * 1e3)
    usd_toe_2014 = g["ely oil products US$m"] / g["ely oil products Mtoe"]
    php_toe = usd_toe_2014 * dubai(year) / dubai(2014) * php_per_usd(year)
    return {"oil-based generation (GWh)": gen_oil_gwh, "toe per MWh (GTAP 2014 / DOE 2014)": toe_mwh,
            "US$ per toe 2014 (GTAP, agents' prices)": usd_toe_2014, f"Dubai {year} / 2014": dubai(year) / dubai(2014),
            f"PHP per USD {year}": php_per_usd(year), f"₱ per toe {year}": php_toe,
            "₱ bn": gen_oil_gwh * 1e3 * toe_mwh * php_toe / 1e9}


def main():
    stamp = dt.datetime.now()
    base = paths.output_dir() / f"electricity_fuel_mix_evidence_{stamp:%Y%m%d_%H%M}"
    s, t = psa_sut(), psa_tt()
    U, SU, names, inds = s["use"], s["supply"], s["products"], s["industries"]
    D = market_shares(SU)
    gap_T = float(np.abs(D @ U - t["T"]).max())
    gap_M = float(np.abs(D @ s["imports"] - t["imports"]).max())
    print(f"1. PSA's I-O = D x Use (fixed product-sales structure, total flows): max gap {gap_T:.1e} (flows), {gap_M:.1e} "
          "(imports), ₱ million")
    p = lambda k: k - 1
    # refining industry's shares of products (what the I-O's 'refined petroleum' row carries)
    ref_sh = pd.Series({names[p(k)]: D[p(REFINED), p(k)] for k in (REFINED, OTHER_PETR, CRUDE_GAS, COAL, ELEC, 42)})
    ref_sh["(other petroleum) crude petroleum and natural gas"] = D[p(OTHER_PETR), p(CRUDE_GAS)]
    print("   Refining industry's share of each product's domestic production (= the share of that product's flows, imports "
          "included, booked in the I-O's refined-petroleum row):")
    print((ref_sh * 100).round(2).to_string())
    # electricity's fuel products: Use table vs I-O rows
    fuels = {COAL: "coal (039)", CRUDE_GAS: "crude petroleum and natural gas (040)", REFINED: "refined petroleum (107)",
             OTHER_PETR: "other petroleum products (108)"}
    go = s["gross_output"][p(ELEC)]
    el = pd.DataFrame({"Use table, products (₱ bn)": {v: U[p(k), p(ELEC)] / 1e3 for k, v in fuels.items()},
                       "I-O Table C.1, industry rows (₱ bn)": {v: t["T"][p(k), p(ELEC)] / 1e3 for k, v in fuels.items()}})
    el.loc["fuels"] = el.sum()
    for c in list(el.columns):
        el[c.replace("(₱ bn)", "(% of output)")] = el[c] / (go / 1e3) * 100
    # decomposition of the I-O's refined-petroleum row into electricity's products
    dec = pd.Series({names[i]: D[p(REFINED), i] * U[i, p(ELEC)] / 1e3 for i in range(240) if D[p(REFINED), i] * U[i, p(ELEC)] > 50})
    print(f"\n2. Electricity industry (240-code 172; output ₱{go / 1e3:,.1f} bn):")
    print(el.round(2).to_string())
    print("   The I-O's refined-petroleum entry for electricity, by the product it really is (₱ bn):")
    print(dec.sort_values(ascending=False).round(2).to_string())
    coal_users = pd.Series(U[p(COAL)] / 1e3, index=inds).sort_values(ascending=False)
    print(f"\n   Coal product (039): supply ₱{(SU[p(COAL)].sum() + s['imports'][p(COAL)]) / 1e3:,.1f} bn (domestic "
          f"{SU[p(COAL)].sum() / 1e3:,.1f}, imports {s['imports'][p(COAL)] / 1e3:,.1f}); intermediate use "
          f"₱{U[p(COAL)].sum() / 1e3:,.1f} bn, of which:")
    print(coal_users.head(6).round(2).to_string())
    ref_sec = pd.Series(SU[:, p(REFINED)] / 1e3, index=names).sort_values(ascending=False)
    el_prod = pd.Series(SU[p(ELEC)] / 1e3, index=inds).sort_values(ascending=False)
    print("   Products of the refining industry (₱ bn):", ref_sec.head(7).round(1).to_dict())
    print(f"   Electricity produced outside the electricity industry: ₱{(SU[p(ELEC)].sum() - SU[p(ELEC), p(ELEC)]) / 1e3:,.1f} bn "
          f"of ₱{SU[p(ELEC)].sum() / 1e3:,.1f} bn;", el_prod.iloc[1:6].round(1).to_dict())
    # DOE, GTAP, prices
    doe = doe_generation()
    g = gtap_2014_power()
    gen = lambda typ, y: float(doe.loc[typ, y])
    print("\n3. DOE gross generation by plant type (GWh; 2003-2020 include off-grid, 2021 on grid only):")
    print(doe[[2014, 2018, 2020, 2021, 2023]].to_string())
    share = doe.loc[["Coal", "Oil-Based", "Natural Gas", "Renewable Energy (RE)"]] / doe.loc["Total Gross Generation"] * 100
    print("   shares (%):", share[[2018, 2023]].round(1).to_dict())
    print("   GTAP 10A 2014, Philippine electricity:", {k: round(v, 3) for k, v in g.items()})
    oil18 = oil_products_bill(gen("Oil-Based", 2018), 2018, g, gen("Oil-Based", 2014))
    offgrid = gen("Oil-Based", 2020) - gen("Oil-Based", 2021)
    oil23 = oil_products_bill(gen("Oil-Based", 2023) + offgrid, 2023, g, gen("Oil-Based", 2014))
    gas18 = U[p(CRUDE_GAS), p(ELEC)] / 1e3
    gas_toe_mwh = g["ely gas Mtoe"] * 1e6 / (gen("Natural Gas", 2014) * 1e3)
    gas_toe18 = gen("Natural Gas", 2018) * 1e3 * gas_toe_mwh
    coal_toe_mwh = g["ely coal Mtoe"] * 1e6 / (gen("Coal", 2014) * 1e3)
    coal_toe18 = gen("Coal", 2018) * 1e3 * coal_toe_mwh
    psa_fuel18 = sum(U[p(k), p(ELEC)] for k in fuels) / 1e3
    coal18_resid = psa_fuel18 - oil18["₱ bn"] - gas18
    # cross-check: PSA coal use x the power sector's share of coal use (GTAP 2014 moved to 2018)
    gr = {"nmm": real_gva_growth("Manufacture of other non-metallic", 2014, 2018),
          "i_s": real_gva_growth("Manufacture of basic metals", 2014, 2018),
          "other": real_gva_growth("Gross Value Added in Manufacturing", 2014, 2018)}
    non_power18 = (g["coal use, nmm (cement etc.) Mtoe"] * gr["nmm"] + g["coal use, i_s (iron and steel) Mtoe"] * gr["i_s"]
                   + g["coal use, other sectors Mtoe"] * gr["other"])
    power_share18 = coal_toe18 / 1e6 / (coal_toe18 / 1e6 + non_power18)
    coal18_share = U[p(COAL)].sum() / 1e3 * power_share18
    usd18 = php_per_usd(2018)
    est = pd.DataFrame({
        "PSA Use table 2018 (₱ bn)": {"coal": U[p(COAL), p(ELEC)] / 1e3, "natural gas (and crude)": gas18,
                                      "oil products": (U[p(REFINED), p(ELEC)] + U[p(OTHER_PETR), p(ELEC)]) / 1e3},
        "PSA I-O 2018, rows (₱ bn)": {"coal": t["T"][p(COAL), p(ELEC)] / 1e3, "natural gas (and crude)": t["T"][p(CRUDE_GAS), p(ELEC)] / 1e3,
                                      "oil products": (t["T"][p(REFINED), p(ELEC)] + t["T"][p(OTHER_PETR), p(ELEC)]) / 1e3},
        "DOE/IEA-consistent 2018 (₱ bn)": {"coal": coal18_resid, "natural gas (and crude)": gas18, "oil products": oil18["₱ bn"]}})
    est.loc["fuels"] = est.sum()
    for c in list(est.columns):
        est[c.replace("(₱ bn)", "(% of output)")] = est[c] / (go / 1e3) * 100
    implied = pd.Series({
        "coal, residual (₱ bn)": coal18_resid, "coal, PSA coal use x power share (₱ bn)": coal18_share,
        "power share of coal use 2018 (GTAP 2014 moved by DOE coal generation and PSA real GVA)": power_share18,
        "PSA's coal booked in basic iron and steel (₱ bn)": U[p(COAL), p(BASIC_IRON_STEEL)] / 1e3,
        "basic iron and steel's share of PSA's intermediate coal": U[p(COAL), p(BASIC_IRON_STEEL)] / U[p(COAL)].sum(),
        "iron and steel's share of coal use, GTAP 2014 (IEA)": g["coal use, i_s (iron and steel) Mtoe"] / g["coal use, all sectors Mtoe"],
        "power's share of coal use, GTAP 2014 (IEA)": g["ely coal Mtoe"] / g["coal use, all sectors Mtoe"],
        "coal burnt by power plants 2018 (Mtoe)": coal_toe18 / 1e6,
        "implied coal price 2018, residual (US$ per GJ)": coal18_resid * 1e9 / usd18 / (coal_toe18 * MJ_PER_TOE / 1e3),
        "implied coal price 2018, residual (US$ per tonne of 5,000 kcal/kg coal)":
            coal18_resid * 1e9 / usd18 / (coal_toe18 * MJ_PER_TOE / 1e3) * GJ_PER_TONNE_5000KCAL,
        "gas burnt 2018 (Mtoe)": gas_toe18 / 1e6,
        "implied gas price 2018, PSA value (US$ per MMBtu)": gas18 * 1e9 / usd18 / (gas_toe18 * MJ_PER_TOE / 1e9 * TBTU_PER_PJ * 1e6),
        "oil: PSA's refined petroleum per kWh of oil-based generation 2018 (₱/kWh)": U[p(REFINED), p(ELEC)] * 1e6 / (gen("Oil-Based", 2018) * 1e6),
        "oil: DOE/IEA-consistent fuel cost per kWh of oil-based generation 2018 (₱/kWh)": oil18["₱ bn"] * 1e9 / (gen("Oil-Based", 2018) * 1e6),
        "oil: PSA's refined petroleum per kilogram of oil equivalent burnt 2018 (₱ per kgoe)":
            U[p(REFINED), p(ELEC)] * 1e6 / (gen("Oil-Based", 2018) * 1e3 * oil18["toe per MWh (GTAP 2014 / DOE 2014)"] * 1e3),
        "Dubai crude 2018 (₱ per kg of crude, about 7.33 bbl per tonne)": dubai(2018) * usd18 * 7.33 / 1e3,
        "oil 2023: off-grid oil-based generation estimate (GWh; DOE 2020 incl. off-grid less 2021 grid only)": offgrid,
        "oil 2023: DOE/IEA-consistent bill (₱ bn)": oil23["₱ bn"]})
    print("\n4. Electricity's fuel bill (240-code 172):")
    print(est.round(2).to_string())
    print(implied.round(4).to_string())
    print("   oil 2018:", {k: round(v, 4) for k, v in oil18.items()})
    print("   oil 2023:", {k: round(v, 4) for k, v in oil23.items()})
    gt = pd.Series({"coal": g["ely coal US$m"], "natural gas (and crude)": g["ely gas US$m"], "oil products": g["ely oil products US$m"]}) / \
        g["ely total cost US$m"] * 100
    cmp_ = pd.DataFrame({"PSA I-O 2018": est["PSA I-O 2018, rows (% of output)"].drop("fuels"),
                         "PSA Use table 2018": est["PSA Use table 2018 (% of output)"].drop("fuels"),
                         "DOE/IEA-consistent 2018": est["DOE/IEA-consistent 2018 (% of output)"].drop("fuels"),
                         "GTAP 10A 2014 (IEA volumes)": gt})
    if ICIO_CSV.exists():
        ic = pd.read_csv(ICIO_CSV, index_col=0)
        m = {"coal (% of output)": "coal", "crude petroleum and natural gas (% of output)": "natural gas (and crude)",
             "refined petroleum and coke (% of output)": "oil products"}
        for c in ("ICIO 2018 (D)", "ICIO 2022 (D)"):
            cmp_[f"OECD {c}"] = ic[c].rename(index=m)
    print("\n5. Fuel cost shares of electricity (% of output or of total cost):")
    print(cmp_.round(2).to_string())
    for name, df in (("io_equals_D_use", pd.Series({"max gap, flows (₱ m)": gap_T, "max gap, imports (₱ m)": gap_M})),
                     ("refining_industry_product_shares", ref_sh), ("electricity_fuels_use_vs_io", el),
                     ("io_refined_entry_for_electricity_by_product", dec), ("coal_users_2018", coal_users),
                     ("doe_generation", doe), ("gtap2014_power", pd.Series(g)), ("fuel_bill_2018", est),
                     ("implied_and_checks", implied), ("oil_2018", pd.Series(oil18)), ("oil_2023", pd.Series(oil23)),
                     ("cost_shares_compared", cmp_)):
        df.to_csv(base.with_name(base.name + f"_{name}.csv"))
    # inputs for tools/build_own_sam_2023.py, version '0.6e'
    rows = []
    for y in (2014, 2018, 2020, 2021, 2023):
        for typ in ("Coal", "Oil-Based", "Natural Gas", "Renewable Energy (RE)", "Total Gross Generation"):
            rows.append((f"DOE gross generation {typ} {y}", gen(typ, y), "GWh",
                         f"DOE 2025 Power Statistics, summary table ({DOE_PDF.name}); 2003-2020 include off-grid, 2021 on grid only"))
    for k in ("ely coal Mtoe", "ely oil products Mtoe", "ely gas Mtoe", "ely coal US$m", "ely oil products US$m", "ely gas US$m",
              "ely total cost US$m", "coal use, all sectors Mtoe", "coal use, nmm (cement etc.) Mtoe",
              "coal use, i_s (iron and steel) Mtoe"):
        rows.append((f"GTAP 10A 2014 Philippines {k}", g[k], "Mtoe" if "Mtoe" in k else "US$ million, agents' prices",
                     "GTAP 10A Data Base 2014 (GTAP10A_GTAP_AY.pkg): gsdvole.har EDF+EIF; basedata.har VDFA+VIFA(+EVFA)"))
    for k, v in (("PSA 2018 Use table: electricity (172) buys coal (039)", U[p(COAL), p(ELEC)] / 1e3),
                 ("PSA 2018 Use table: electricity (172) buys crude petroleum and natural gas (040)", gas18),
                 ("PSA 2018 Use table: electricity (172) buys refined petroleum (107)", U[p(REFINED), p(ELEC)] / 1e3),
                 ("PSA 2018 Use table: electricity (172) buys other petroleum products (108)", U[p(OTHER_PETR), p(ELEC)] / 1e3),
                 ("PSA 2018 Use table: electricity (172) gross output", go / 1e3),
                 ("PSA 2018 Use table: coal (039) intermediate use, all industries", U[p(COAL)].sum() / 1e3),
                 ("PSA 2018 Use table: basic iron and steel (123) buys coal (039)", U[p(COAL), p(BASIC_IRON_STEEL)] / 1e3)):
        rows.append((k, v, "₱ bn", f"PSA 2018 Supply and Use Tables ({SUT.name}), 240_Use"))
    inp = pd.DataFrame(rows, columns=["key", "value", "unit", "source"])
    head = ("# Inputs for tools/build_own_sam_2023.py version '0.6e' (electricity's fuel mix). RESEARCH, NOT FOR "
            "CITATION.\n# Written by research/electricity_fuel_mix/fuel_mix_evidence.py on "
            f"{stamp:%Y-%m-%d %H:%M} from files on disk only; each row names its source. Prices (Dubai, PHP per USD) are "
            "read by the builder from data/raw/macro.\n")
    with open(INPUTS_CSV, "w", encoding="utf-8", newline="") as f:
        f.write(head)
        inp.to_csv(f, index=False)
    print("\nsaved", base.name + "_*.csv in", base.parent, "and", INPUTS_CSV)


if __name__ == "__main__":
    main()
